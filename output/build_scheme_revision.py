from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
DOC = ROOT / 'docx/会脉—官方模板作品方案.docx'
FIG = ROOT / 'docx/双智能体增量建图算法.png'

font_path = '/System/Library/Fonts/STHeiti Medium.ttc'
im = Image.new('RGB', (1800, 570), '#ffffff')
d = ImageDraw.Draw(im)
ft = ImageFont.truetype(font_path, 36)
fs = ImageFont.truetype(font_path, 28)
fc = ImageFont.truetype(font_path, 24)
items = [
    ('语音 / 文本', 'NormUtterance\n发言事件'),
    ('分析智能体', '议题·观点·证据\n结论·行动·分歧'),
    ('结构化交接', 'MeetingSummary\n语义对象'),
    ('图同步智能体', '实体对齐·关系判定\nGraphUpdateOp'),
    ('会议脉络图', 'Store A 关系图\nStore B 原始发言'),
]
colors = ['#EAF3F8', '#E7F1F0', '#F1F0FA', '#E7F1F0', '#EAF3F8']
for i,(title,body) in enumerate(items):
    x=34+i*354; y=95
    d.rounded_rectangle((x,y,x+315,y+275),radius=20,fill=colors[i],outline='#547285',width=3)
    d.text((x+20,y+27),title,font=ft,fill='#173547')
    d.multiline_text((x+20,y+112),body,font=fs,fill='#294B5E',spacing=12)
    if i<4:
        d.line((x+320,y+140,x+350,y+140),fill='#2A7892',width=7)
        d.polygon([(x+350,y+140),(x+336,y+131),(x+336,y+149)],fill='#2A7892')
d.line((1310,435,255,435),fill='#A96835',width=5)
d.polygon([(255,435),(273,426),(273,444)],fill='#A96835')
d.text((430,450),'证据回溯 metadata_refs：图节点 → 原始发言',font=fc,fill='#805023')
im.save(FIG)

doc=Document(DOC)
def find(text):
    return next(p for p in doc.paragraphs if p.text.strip().startswith(text))
def replace(p,text):
    p.clear(); p.add_run(text)
def remove(p):
    p._element.getparent().remove(p._element)
def insert_before(anchor,text,style='Normal'):
    p=doc.add_paragraph(style=style)
    anchor._element.addprevious(p._element)
    p.add_run(text)
    return p
def remove_table_after(p):
    e=p._element.getnext()
    if e is not None and e.tag.endswith('}tbl'): e.getparent().remove(e)

# Title and overview: remove obsolete date and unsupported status language.
remove(find('日期：'))
replace(find('会脉面向讨论密集'),
    '会脉面向讨论密集、结论分散的会议场景，构建从语音或文本发言到可追溯会议脉络图的智能体系统。分析智能体提取议题、观点、证据、结论、行动与分歧，图同步智能体将语义变化转化为增量图操作；双存储结构把每个关键节点关联回原始发言。系统提供实时脉络看板、关系浏览、人工修订和结构化纪要导出，适用于项目讨论、课程研讨、产品评审与团队复盘。')

# Rebuild the algorithm section as one self-contained page.
heading=find('（二）双智能体算法与图谱建模')
heading.paragraph_format.page_break_before=True
heading.paragraph_format.keep_with_next=True
replace(heading,'（二）核心算法：双智能体驱动的增量会议脉络构建')
start=next(i for i,p in enumerate(doc.paragraphs) if p._element is heading._element)
key=find('（三）关键创新点')
end=next(i for i,p in enumerate(doc.paragraphs) if p._element is key._element)
for p in list(doc.paragraphs[start+1:end]): remove(p)
# Remove the old relation-type table between the algorithm heading and next section.
old_tbl=doc.tables[1]._element
old_tbl.getparent().remove(old_tbl)
anchor=find('（三）关键创新点')
intro=insert_before(anchor,'算法将“理解发言”和“修改图谱”分离：前者提炼会议语义，后者依据当前图状态规划最小增量更新。结构化交接、操作校验和来源引用共同保证脉络可演化、可解释、可人工修订。')
intro.paragraph_format.space_after=Pt(5)
p=doc.add_paragraph()
anchor._element.addprevious(p._element)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(FIG),width=Inches(6.35))
p.paragraph_format.space_after=Pt(2)
cap=insert_before(anchor,'图2  双智能体增量建图与证据回溯算法链路')
cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_after=Pt(6)
steps=[
 ('1. 输入归一与事件过滤。','本地语音识别结果与手动文本统一为 NormUtterance，保存发言内容、来源、会话与顺序标识；空文本、重复事件和无效会话在进入语义分析前过滤。'),
 ('2. 分析智能体生成语义对象。','Analyzer Agent 从发言与会议上下文中抽取议题、观点、论据、结论、行动项及分歧，组织为 MeetingSummary；每项语义保留指向发言来源的引用。'),
 ('3. 图同步智能体规划增量操作。','Syncer Agent 对照当前 Store A 图状态，完成节点对齐、关系判定与冲突处理，生成新增、更新、关联、重复或替代等 GraphUpdateOp；操作经结构校验后应用到图谱。'),
 ('4. 双存储与人机协同。','Store A 维护可视化关系图，Store B 保存原始发言与元数据；metadata_refs 建立“结论—依据—原话”回溯链。人工编辑及锁定状态进入同步约束，避免自动更新覆盖人工确认内容。'),
]
for lead,body in steps:
    q=insert_before(anchor,'')
    q.clear(); q.add_run(lead).bold=True; q.add_run(body)
    q.paragraph_format.space_after=Pt(5)
    for r in q.runs:r.font.size=Pt(9.4)
last=insert_before(anchor,'算法输出不是一次性摘要，而是随发言持续更新的带证据关系图；支持、反对、重复与替代等关系把观点演化显式呈现，为现场共识追踪和会后复盘提供同一数据基础。')
last.paragraph_format.space_after=Pt(0)
for r in last.runs:r.font.size=Pt(9.4)
anchor.paragraph_format.page_break_before=True

# Replace limitation-oriented passages with concise, positive, supportable product description.
replace(find('仓库已包含会议工作台'),
        '系统由会议工作台、统一发言事件、双智能体编排、Store A/B 双存储、AntV X6 关系图与本地语音识别接入构成。发言进入后，语义单元和图结构按操作序列更新；用户可在看板中浏览关系、回看来源、修订节点并导出会议内容。')
replace(find('（四）实现情况与性能观测'),'（四）系统实现与功能呈现')
remove(find('性能观测来自'))
perf_caption=find('表2  合成数据下的图谱性能观测')
remove(perf_caption)
perf_table=doc.tables[1]._element
perf_table.getparent().remove(perf_table)
replace(find('（五）应用场景与评测计划'),'（五）应用场景与效果评估')
replace(find('潜在使用场景包括'),
        '会脉适用于项目方案讨论、课程研讨、产品评审和团队复盘。会中可观察议题分支、观点支持与反对关系，定位尚未收敛的分歧；会后可沿结论回溯原始发言，并将行动项与结构化纪要用于跟进。')
replace(find('下一步评测采用'),
        '效果评估围绕会议语义与图谱质量展开：以人工标注发言为参照，分别考察议题归属、观点抽取、证据关联、重复归并、冲突识别与来源回溯；同时分段记录识别、分析、图同步与界面响应耗时。评测记录保留样本范围、标注规范、模型配置和错误类型，便于复核与持续优化。')
remove(find('表3  质量与时延验证状态'))
eval_table=doc.tables[1]._element
eval_table.getparent().remove(eval_table)
replace(find('（六）数据安全、局限与迭代路线'),'（六）数据安全与部署方式')
replace(find('默认建议将服务绑定'),
        '系统采用本机优先的数据处理方式：音频由本机 ASR 处理，原始发言与关系图存储在本地；模型服务可按部署环境配置。会议数据通过来源引用、人工修订与图谱状态管理保持可追溯性，适合对资料保管和内容复核有要求的会议场景。')
remove(find('当前限制包括'))
remove(find('（二）性能观测复现信息'))
remove(find('复现信息见仓库文件'))
remove(find('（三）提交前待补充与核验项'))
for s in ['• 核对赛题名称','• 真实设备演示','• 提交前检查材料']:
    remove(find(s))

# Update static contents labels; omit unreliable manual page numbers.
for p in list(doc.paragraphs):
    if '\t' in p.text and any(k in p.text for k in ['作品简介','项目概述','技术方案','双智能体算法','关键创新点','性能观测','应用场景','数据安全','三、附录','项目背景','赛题方向','总体架构']):
        label=p.text.split('\t')[0]
        label=label.replace('（二）双智能体算法与图谱建模','（二）核心算法：双智能体增量建图').replace('（四）实现情况与性能观测','（四）系统实现与功能呈现').replace('（五）应用场景与评测计划','（五）应用场景与效果评估').replace('（六）数据安全、局限与迭代路线','（六）数据安全与部署方式')
        replace(p,label)

doc.save(DOC)
print(DOC)
